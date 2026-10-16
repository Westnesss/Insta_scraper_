"""
Genera resumen.xlsx a partir de los JSON de la carpeta datos/.

Hojas:
  - Perfiles:       una fila por perfil (usuario, seguidores, seguidos, descripción...)
  - Publicaciones:  una fila por publicación de todos los perfiles

Uso:  pip install openpyxl   y luego   python exportar_excel.py
Los JSON siguen siendo la fuente de datos; el Excel es solo una vista para leer y comparar.
(No incluye las listas de seguidores/seguidos a propósito.)
"""
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CARPETA_DATOS = "datos"
SALIDA = "resumen.xlsx"


def leer_json(ruta):
    if not os.path.exists(ruta):
        return None
    with open(ruta, encoding="utf-8") as f:
        return json.load(f)


def a_numero(texto):
    """'4.000' -> 4000, '98 M' -> 98000000, '1,2 M' -> 1200000. Si no se entiende, devuelve el texto."""
    if texto is None:
        return None
    t = str(texto).strip().replace("\u00a0", " ")
    m = re.fullmatch(r"([\d.,]+)\s?([KkMm])", t)
    if m:
        base = float(m.group(1).replace(",", "."))
        return int(round(base * (1_000 if m.group(2).lower() == "k" else 1_000_000)))
    if re.fullmatch(r"[\d.,]+", t):
        return int(t.replace(".", "").replace(",", ""))
    return texto


def dar_formato(hoja, anchos):
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    hoja.freeze_panes = "A2"
    for col, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")


def archivos_de(pub):
    if "imagenes" in pub:   # formato actual
        return [i["archivo"] for i in pub["imagenes"]]
    return [pub["imagen_local"]] if pub.get("imagen_local") else []   # formato antiguo


def main():
    if not os.path.isdir(CARPETA_DATOS):
        print(f"No existe la carpeta '{CARPETA_DATOS}/'. Ejecuta primero el scraper.")
        return

    wb = Workbook()
    hoja_perfiles = wb.active
    hoja_perfiles.title = "Perfiles"
    hoja_perfiles.append(["Usuario", "Seguidores", "Seguidos", "Publicaciones (total)",
                          "Descripción del perfil", "URL", "Publicaciones extraídas"])
    hoja_pubs = wb.create_sheet("Publicaciones")
    hoja_pubs.append(["Usuario", "N°", "URL", "Descripción", "Resumen", "N° de imágenes", "Archivos de imagen"])

    for usuario in sorted(os.listdir(CARPETA_DATOS)):
        carpeta = os.path.join(CARPETA_DATOS, usuario)
        if not os.path.isdir(carpeta):
            continue
        perfil = leer_json(os.path.join(carpeta, "perfil.json")) or {}
        pubs = leer_json(os.path.join(carpeta, "publicaciones.json")) or []

        hoja_perfiles.append([
            usuario,
            a_numero(perfil.get("seguidores")),
            a_numero(perfil.get("seguidos")),
            a_numero(perfil.get("publicaciones_total")),
            perfil.get("descripcion_perfil", ""),
            perfil.get("url", f"https://www.instagram.com/{usuario}/"),
            len(pubs),
        ])
        for pub in pubs:
            archivos = archivos_de(pub)
            hoja_pubs.append([
                usuario, pub.get("id"), pub.get("url"), pub.get("descripcion"),
                pub.get("resumen"), len(archivos),
                "\n".join(f"{CARPETA_DATOS}/{usuario}/{a}" for a in archivos),
            ])

    dar_formato(hoja_perfiles, [22, 14, 12, 20, 50, 38, 22])
    dar_formato(hoja_pubs, [22, 6, 38, 60, 45, 14, 45])
    wb.save(SALIDA)
    print(f"Listo: {SALIDA} ({hoja_perfiles.max_row - 1} perfiles, {hoja_pubs.max_row - 1} publicaciones)")


if __name__ == "__main__":
    main()